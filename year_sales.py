import openpyxl
import pandas as pd
from pandas import read_excel
from openpyxl import Workbook
from openpyxl.chart import LineChart ,Reference

wb=openpyxl.Workbook
# wb=pd.read_excel('C://Users/user794/Documents/sales_monthwise.xlsx',sheet_name='SALES_JAN')
# print(exldata1)
wb=openpyxl.load_workbook('C://Users/user794/Documents/sales_monthwise.xlsx')

sheet1=wb['SALES_JAN']
sheet2=wb['SALES_FEB']
# sheet['A1:A6']

chart=LineChart()
data1=Reference(worksheet=sheet1,min_row=1,max_row=6,min_col=1
,max_col=1 )
chart.add_data(data1,titles_from_data=True)
sheet1.add_chart(chart,"E2")

###############3
chart=LineChart()
data2=Reference(worksheet=sheet2,min_row=1,max_row=6,min_col=1
,max_col=1 )
chart.add_data(data2,titles_from_data=True)
sheet2.add_chart(chart,"G2")
df1=pd.read_excel('C://Users/user794/Documents/sales_monthwise.xlsx',sheet_name='SALES_JAN')
df2=pd.read_excel('C://Users/user794/Documents/sales_monthwise.xlsx',sheet_name='SALES_FEB')

# print(df1)


wb.save('C://Users/user794/Documents/sales_monthwise1.xlsx')

